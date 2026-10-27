"""06: Cek silang Excel (Q13). Untuk 3 operator (satu per stasiun), time-to-target dihitung
dengan formula Excel dari hari valid, lalu dihitung ulang oleh LibreOffice (headless).
Hasil dibandingkan dengan views SQL. Workbook dengan nilai terhitung: outputs/crosscheck.xlsx.
"""
import shutil
import subprocess
import sys
import tempfile
from pathlib import Path

sys.path.insert(0, str(Path(__file__).resolve().parents[2]))
import openpyxl
from openpyxl.styles import Font
import pandas as pd
import common as C

CHECK_OPS = ["OP-01", "OP-05", "OP-08"]              # ST-A, ST-B, ST-C
days = pd.read_csv(C.RES2 / "valid_days_sql.csv")
sql = pd.read_csv(C.RES2 / "ttt_per_operator.csv").set_index("operator_code")

wb = openpyxl.Workbook()
wb.remove(wb.active)
summary = wb.create_sheet("RINGKASAN")
summary.append(["operator", "stasiun", "TTT Excel", "TTT SQL", "cocok"])
for c in summary[1]:
    c.font = Font(bold=True)

for op in CHECK_OPS:
    g = days[days.operator_code == op].sort_values(["valid_day", "work_day"]).reset_index(drop=True)
    ws = wb.create_sheet(op)
    ws.append(["hari_valid", "hari_kerja", "ok_qty", "target", "persen_target", "memenuhi", "streak_awal"])
    for c in ws[1]:
        c.font = Font(bold=True)
    n = len(g)
    for i, r in g.iterrows():
        row = i + 2
        ws.append([int(r.valid_day), int(r.work_day), int(r.ok_qty), int(r.target),
                   f"=C{row}/D{row}", f"=IF(C{row}>=D{row},1,0)",
                   # hari pertama dari 3 hari berturut-turut memenuhi (dua baris di bawahnya)
                   (f"=IF(AND(F{row}=1,F{row + 1}=1,F{row + 2}=1),1,0)" if i + 2 < n else 0)])
    last = n + 1
    ws["I1"] = "TTT (hari valid)"
    ws["I2"] = f'=IFERROR(INDEX(A2:A{last},MATCH(1,G2:G{last},0)),"tidak tercapai")'
    ws["I1"].font = Font(bold=True)
    ws.column_dimensions["E"].width = 14
    ws.column_dimensions["I"].width = 20
    ttt_row = len(summary["A"])
    stn = sql.loc[op, "station_code"]
    summary.append([op, stn, f"='{op}'!I2", int(sql.loc[op, "ttt_day"]), None])
    summary.cell(row=ttt_row + 1, column=5, value=f"=IF(C{ttt_row + 1}=D{ttt_row + 1},\"ya\",\"tidak\")")

tmp = Path(tempfile.mkdtemp())
src = tmp / "crosscheck_raw.xlsx"
wb.save(src)
# LibreOffice menghitung ulang formula, lalu hasilnya disimpan sebagai xlsx dengan nilai
subprocess.run(["soffice", "--headless", "--calc", "--convert-to", "xlsx", "--outdir", str(tmp / "out"), str(src)],
               check=True, capture_output=True, timeout=240)
calc = tmp / "out" / src.name
shutil.copy(calc, C.OUT_XLSX / "crosscheck.xlsx")

vals = openpyxl.load_workbook(calc, data_only=True)
rows = []
for op in CHECK_OPS:
    x = vals[op]["I2"].value
    rows.append(dict(operator_code=op, ttt_excel=x, ttt_sql=int(sql.loc[op, "ttt_day"])))
res = pd.DataFrame(rows)
res["cocok"] = (pd.to_numeric(res.ttt_excel, errors="coerce") == res.ttt_sql).astype(int)
res.to_csv(C.RES2 / "crosscheck_excel.csv", index=False)
print(res.to_string(index=False))
assert res.cocok.sum() == len(res), "Excel dan SQL tidak cocok"
